import html
import os

from flask import (
    Flask,
    render_template,
    request,
    redirect,
    jsonify,
    flash,
    url_for,
    current_app,
    send_file,
    send_from_directory,
    abort,
)
from markupsafe import Markup
from config import Config
from models import db, User, School, Student, Request, Movement, Notification
from flask_login import LoginManager, login_user, login_required, logout_user, current_user
from datetime import datetime
from sqlalchemy import inspect, text, func

from curriculum_ug import (
    UG_CLASS_CODES,
    UG_PRIMARY_LEVELS,
    UG_SECONDARY_LEVELS,
    label_for_code,
)
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from excel_import import (
    build_learner_import_template,
    filter_duplicate_rfids_in_file,
    parse_learner_rows,
)
from auth_utils import (
    MIN_PASSWORD_LENGTH,
    hash_password,
    password_matches,
    upgrade_password_if_needed,
)
from parent_accounts import resolve_or_create_parent
from parent_flash import flash_new_parent_credentials
from parent_credentials import (
    build_parent_logins_workbook,
    flash_parent_logins_download,
    get_import_parent_logins,
)
from learner_ids import (
    assign_student_code,
    backfill_all_student_codes,
    derive_school_initials,
    display_student_id,
    find_student_by_lookup,
)
from student_photos import (
    can_view_student_photo,
    delete_student_photo_file,
    save_student_photo_upload,
    student_photo_dir,
    student_photo_url,
)
from notify import (
    KIND_DEPARTURE,
    KIND_EXIT_REQUEST,
    KIND_REQUEST_DECISION,
    KIND_RETURN,
    KIND_SCHOOL_PENDING,
    notify_superadmins,
    alert_user,
    class_display,
    notify_class_teacher,
    notify_school_admins,
    notify_gate_movement,
    notify_student_stakeholders,
)

app = Flask(__name__)
app.config.from_object(Config)

db.init_app(app)

login_manager = LoginManager(app)
login_manager.login_view = "login"


def _safe_add_column(table: str, col_name: str, col_def: str):
    """SQLite: add a column if missing. Silently skips if already present."""
    try:
        insp = inspect(db.engine)
        if table not in insp.get_table_names():
            return
        cols = {c["name"] for c in insp.get_columns(table)}
        if col_name in cols:
            return
        with db.engine.begin() as conn:
            conn.execute(text(f"ALTER TABLE {table} ADD COLUMN {col_name} {col_def}"))
    except Exception:
        pass


def run_migrations():
    _safe_add_column("student", "class_level", "VARCHAR(10) NOT NULL DEFAULT ''")
    _safe_add_column("user", "class_level", "VARCHAR(10)")
    _safe_add_column("request", "created_at", "DATETIME")
    _safe_add_column("user", "stream", "VARCHAR(50)")
    _safe_add_column("student", "stream", "VARCHAR(50)")
    _safe_add_column("notification", "kind", "VARCHAR(30)")
    _safe_add_column("student", "photo_filename", "VARCHAR(120)")
    _safe_add_column("school", "initials", "VARCHAR(3)")
    _safe_add_column("student", "student_code", "VARCHAR(10)")
    if not current_app.config.get("TESTING"):
        try:
            if Student.query.filter(Student.student_code.is_(None)).first():
                backfill_all_student_codes()
        except Exception:
            db.session.rollback()


def ensure_superadmin_exists():
    """Seed superadmin if missing (works with `flask run`, not only `python app.py`)."""
    if current_app.config.get("TESTING"):
        return
    email = "superadmin@system.com"
    if User.query.filter(func.lower(User.email) == email.lower()).first():
        return
    db.session.add(
        User(
            name="Super Admin",
            email=email,
            password=hash_password("admin123"),
            role="superadmin",
            school_id=None,
        )
    )
    db.session.commit()


@app.before_request
def _app_init_once():
    if current_app.config.get("_app_init_done"):
        return
    db.create_all()
    run_migrations()
    ensure_superadmin_exists()
    current_app.config["_app_init_done"] = True


@app.template_filter("student_display_id")
def student_display_id_filter(student):
    return display_student_id(student)


@app.context_processor
def inject_globals():
    ctx = {
        "nav_school": None,
        "unread_notif_count": 0,
        "latest_notif_id": 0,
        "live_popup_role": False,
        "student_photo_url": student_photo_url,
        "display_student_id": display_student_id,
    }
    if getattr(current_user, "is_authenticated", False):
        if getattr(current_user, "school_id", None):
            ctx["nav_school"] = School.query.get(current_user.school_id)
        ctx["unread_notif_count"] = Notification.query.filter_by(
            user_id=current_user.id, is_read=False
        ).count()
        latest = (
            db.session.query(func.max(Notification.id))
            .filter_by(user_id=current_user.id)
            .scalar()
        )
        ctx["latest_notif_id"] = latest or 0
        ctx["live_popup_role"] = current_user.role in (
            "admin",
            "teacher",
            "parent",
            "superadmin",
        )
    else:
        ctx["latest_notif_id"] = 0
        ctx["live_popup_role"] = False
    return ctx


@login_manager.user_loader
def load_user(user_id):
    return User.query.get(int(user_id))


# =======================
# HOME (LANDING PAGE)
# =======================
@app.route('/')
def home():
    if current_user.is_authenticated:
        return redirect(url_for("dashboard"))
    return render_template("home.html")


# =======================
# LOGIN
# =======================
@app.route('/login', methods=['GET', 'POST'])
def login():
    if current_user.is_authenticated:
        return redirect(url_for("dashboard"))

    if request.method == 'POST':
        email = (
            request.form.get("work_email") or request.form.get("email") or ""
        ).strip()
        password = request.form.get("auth_pass") or request.form.get("password") or ""

        if not email:
            return render_template("login.html", error="Email is required")

        user = User.query.filter(func.lower(User.email) == email.lower()).first()

        if not user:
            return render_template(
                "login.html",
                error="No account for that email. Check spelling or register your school.",
            )

        if not password_matches(user.password, password):
            return render_template('login.html', error="Wrong password")

        if upgrade_password_if_needed(user, password):
            db.session.commit()

        # Check school approval
        if user.school_id:
            school = School.query.get(user.school_id)
            if school and school.status != 'approved':
                return render_template('login.html', error="School not approved yet")

        login_user(user)
        return redirect('/dashboard')

    return render_template('login.html')


# =======================
# REGISTER SCHOOL + ADMIN
# =======================
@app.route("/register_school", methods=["GET", "POST"])
def register_school():
    if current_user.is_authenticated:
        flash(
            "You are already signed in. Sign out first if you need to register another institution.",
            "error",
        )
        return redirect(url_for("dashboard"))

    if request.method == "POST":
        school_name = request.form.get('school_name')
        admin_name = request.form.get('admin_name')
        admin_email = (request.form.get("admin_email") or "").strip().lower()
        admin_password = request.form.get('admin_password')

        if not all([school_name, admin_name, admin_email, admin_password]):
            flash("All fields are required.", "error")
            return redirect(url_for("register_school"))

        existing = User.query.filter_by(email=admin_email).first()
        if existing:
            flash("An account with this admin email already exists.", "error")
            return redirect(url_for("register_school"))

        school = School(
            name=school_name,
            status="pending",
            initials=derive_school_initials(school_name),
        )
        db.session.add(school)
        db.session.flush()

        admin = User(
            name=admin_name,
            email=admin_email,
            password=hash_password(admin_password),
            role='admin',
            school_id=school.id
        )

        db.session.add(admin)
        db.session.flush()

        notify_superadmins(
            f'New school "{school.name}" registered and awaits approval. '
            f"Administrator: {admin_name} ({admin_email}).",
            email_subject="CampusPass: new school awaiting approval",
        )
        db.session.commit()

        flash(
            "School registered. After a super administrator approves your campus, sign in on the home page with the admin email and password you just created.",
            "success",
        )
        return redirect(url_for("login"))

    return render_template("register_school.html")


# =======================
# ADMIN MANAGES STAFF
# =======================
STAFF_ROLES = ("nurse", "teacher", "guard", "parent")
SCHOOL_STAFF_ROLES = ("nurse", "teacher", "guard")
TEAM_VIEWS = ("staff", "parents", "create")


def school_staff_query(school_id):
    return (
        User.query.filter(
            User.school_id == school_id,
            User.role.in_(STAFF_ROLES),
        )
        .order_by(User.role, User.name)
    )


def school_staff_members_query(school_id):
    return (
        User.query.filter(
            User.school_id == school_id,
            User.role.in_(SCHOOL_STAFF_ROLES),
        )
        .order_by(User.role, User.name)
    )


def school_parents_query(school_id):
    return (
        User.query.filter_by(school_id=school_id, role="parent")
        .order_by(User.name)
    )


def team_view_from_request():
    view = (request.args.get("view") or "staff").strip().lower()
    return view if view in TEAM_VIEWS else "staff"


def admin_can_manage_staff(admin, target):
    if not admin or admin.role != "admin" or not admin.school_id:
        return False
    if not target or target.school_id != admin.school_id:
        return False
    if target.role not in STAFF_ROLES:
        return False
    if target.id == admin.id:
        return False
    return True


@app.route('/create_user', methods=['GET', 'POST'])
@login_required
def create_user():

    if current_user.role != "admin":
        flash("You do not have permission to manage team accounts.", "error")
        return redirect(url_for("dashboard"))

    if request.method == "POST":
        name = request.form.get("name")
        email = (request.form.get("email") or "").strip().lower()
        password = request.form.get("password")
        role = request.form.get("role")

        allowed_roles = ["nurse", "teacher", "guard", "parent"]

        if role not in allowed_roles:
            flash("Invalid role selected.", "error")
            return redirect(url_for("create_user", view="create"))

        existing = User.query.filter_by(email=email).first()
        if existing:
            flash("This email is already in use.", "error")
            return redirect(url_for("create_user", view="create"))

        teacher_class = ""
        teacher_stream = ""
        if role == "teacher":
            teacher_class = (request.form.get("class_level") or "").strip().upper()
            teacher_stream = (request.form.get("stream") or "").strip()
            if teacher_class not in UG_CLASS_CODES:
                flash("Teachers must be assigned a valid class (P1–P7 or S1–S6).", "error")
                return redirect(url_for("create_user", view="create"))

        user = User(
            name=name,
            email=email,
            password=hash_password(password),
            role=role,
            school_id=current_user.school_id,
            class_level=teacher_class if role == "teacher" else None,
            stream=teacher_stream if role == "teacher" and teacher_stream else None,
        )

        db.session.add(user)
        db.session.commit()

        extra = f" assigned to {class_display(teacher_class, teacher_stream)}" if teacher_class else ""
        flash(f"Account created for {name}{extra}. They can sign in with the email and password you set.", "success")
        return redirect(
            url_for("create_user", view="parents" if role == "parent" else "staff")
        )

    view = team_view_from_request()
    staff_members = school_staff_members_query(current_user.school_id).all()
    parent_accounts = school_parents_query(current_user.school_id).all()
    parent_ids = [p.id for p in parent_accounts]
    parent_child_counts = {}
    if parent_ids:
        rows = (
            db.session.query(Student.parent_id, func.count(Student.id))
            .filter(Student.parent_id.in_(parent_ids))
            .group_by(Student.parent_id)
            .all()
        )
        parent_child_counts = {pid: cnt for pid, cnt in rows}

    return render_template(
        "create_user.html",
        view=view,
        staff_members=staff_members,
        parent_accounts=parent_accounts,
        parent_child_counts=parent_child_counts,
        ug_primary_levels=UG_PRIMARY_LEVELS,
        ug_secondary_levels=UG_SECONDARY_LEVELS,
    )


@app.route("/staff/<int:user_id>/edit", methods=["GET", "POST"])
@login_required
def edit_staff(user_id):
    if current_user.role != "admin":
        flash("You do not have permission to manage team accounts.", "error")
        return redirect(url_for("dashboard"))

    member = User.query.get_or_404(user_id)
    if not admin_can_manage_staff(current_user, member):
        flash("You cannot edit this account.", "error")
        return redirect(
            url_for(
                "create_user",
                view="parents" if member.role == "parent" else "staff",
            )
        )

    if request.method == "POST":
        name = (request.form.get("name") or "").strip()
        email = (request.form.get("email") or "").strip().lower()
        role = request.form.get("role")
        new_password = request.form.get("password") or ""

        if role not in STAFF_ROLES:
            flash("Invalid role selected.", "error")
            return redirect(url_for("edit_staff", user_id=user_id))

        if not name or not email:
            flash("Name and email are required.", "error")
            return redirect(url_for("edit_staff", user_id=user_id))

        clash = User.query.filter(
            func.lower(User.email) == email,
            User.id != member.id,
        ).first()
        if clash:
            flash("This email is already in use by another account.", "error")
            return redirect(url_for("edit_staff", user_id=user_id))

        teacher_class = None
        teacher_stream = None
        if role == "teacher":
            teacher_class = (request.form.get("class_level") or "").strip().upper()
            teacher_stream = (request.form.get("stream") or "").strip()
            if teacher_class not in UG_CLASS_CODES:
                flash("Teachers must be assigned a valid class (P1–P7 or S1–S6).", "error")
                return redirect(url_for("edit_staff", user_id=user_id))
        elif member.role == "teacher" and role != "teacher":
            teacher_class = None
            teacher_stream = None

        member.name = name
        member.email = email
        member.role = role
        member.class_level = teacher_class if role == "teacher" else None
        member.stream = (
            teacher_stream if role == "teacher" and teacher_stream else None
        )

        if new_password.strip():
            if len(new_password) < MIN_PASSWORD_LENGTH:
                flash(
                    f"Password must be at least {MIN_PASSWORD_LENGTH} characters.",
                    "error",
                )
                return redirect(url_for("edit_staff", user_id=user_id))
            member.password = hash_password(new_password)

        db.session.commit()
        extra = (
            f" ({class_display(member.class_level, member.stream)})"
            if member.role == "teacher" and member.class_level
            else ""
        )
        flash(f"Updated account for {member.name}{extra}.", "success")
        return redirect(
            url_for(
                "create_user",
                view="parents" if member.role == "parent" else "staff",
            )
        )

    return render_template(
        "staff_edit.html",
        member=member,
        list_view="parents" if member.role == "parent" else "staff",
        ug_primary_levels=UG_PRIMARY_LEVELS,
        ug_secondary_levels=UG_SECONDARY_LEVELS,
    )


@app.route("/staff/<int:user_id>/delete", methods=["POST"])
@login_required
def delete_staff(user_id):
    if current_user.role != "admin":
        flash("You do not have permission to manage team accounts.", "error")
        return redirect(url_for("dashboard"))

    member = User.query.get_or_404(user_id)
    if not admin_can_manage_staff(current_user, member):
        flash("You cannot delete this account.", "error")
        return redirect(url_for("create_user", view=team_view_from_request()))

    name = member.name or member.email
    role = member.role
    redirect_view = "parents" if role == "parent" else "staff"

    if role == "parent":
        linked = Student.query.filter_by(
            parent_id=member.id, school_id=current_user.school_id
        ).count()
        if linked:
            Student.query.filter_by(parent_id=member.id).update(
                {Student.parent_id: None}, synchronize_session=False
            )
            flash(
                f"{name} was removed. {linked} learner(s) were unlinked from this parent account.",
                "warning",
            )
        else:
            flash(f"Parent account for {name} was deleted.", "success")
    else:
        flash(f"{name} ({role}) was removed from your school.", "success")

    Notification.query.filter_by(user_id=member.id).delete(
        synchronize_session=False
    )
    db.session.delete(member)
    db.session.commit()
    return redirect(url_for("create_user", view=redirect_view))


def school_scoped_requests(user):
    """Requests visible to the current user."""
    q = Request.query.order_by(Request.id.desc())
    if user.role == "superadmin":
        return q.all()
    if user.role == "parent":
        child_ids = [
            s.id for s in Student.query.filter_by(parent_id=user.id).all()
        ]
        if not child_ids:
            return []
        return q.filter(Request.student_id.in_(child_ids)).all()
    if not user.school_id:
        return []
    ids = [s.id for s in Student.query.filter_by(school_id=user.school_id).all()]
    if not ids:
        return []
    return q.filter(Request.student_id.in_(ids)).all()


def roster_for_requester(user):
    """Learners the user may pick when creating a movement request."""
    if user.role == "parent":
        return (
            Student.query.filter_by(parent_id=user.id)
            .order_by(Student.class_level, Student.name)
            .all()
        )
    if user.school_id and user.role in ("nurse", "teacher"):
        q = Student.query.filter_by(school_id=user.school_id)
        if user.role == "teacher" and user.class_level:
            q = q.filter_by(class_level=user.class_level)
            if user.stream:
                q = q.filter_by(stream=user.stream)
        return q.order_by(Student.class_level, Student.stream, Student.name).all()
    return []


def students_picker_payload(students):
    """JSON-friendly roster rows for the nurse learner search picker."""
    return [
        {
            "id": s.id,
            "name": s.name or "",
            "class_level": s.class_level or "",
            "stream": s.stream or "",
            "rfid": s.rfid_tag or "",
            "student_code": s.student_code or "",
            "label": (
                f"{s.name} — {class_display(s.class_level, s.stream)}"
                + (f" · {s.student_code}" if s.student_code else "")
            ),
        }
        for s in students
    ]


def can_manage_student(user, student):
    if not student or not user.school_id or student.school_id != user.school_id:
        return False
    if user.role == "admin":
        return True
    if user.role == "teacher" and user.class_level:
        if student.class_level != user.class_level:
            return False
        if user.stream and student.stream != user.stream:
            return False
        return True
    return False


def link_parent_for_learner(school_id, student_name: str, keep_parent_id=None):
    """
    Resolve parent from form: parent_email creates/links account; parent_id links existing.
    Returns (parent_id, error_message, new_credentials list).
    """
    parent_email = (request.form.get("parent_email") or "").strip()
    parent_name = (request.form.get("parent_name") or "").strip()
    parent_raw = (request.form.get("parent_id") or "").strip()
    new_creds: list[tuple[str, str, str]] = []

    if not parent_email and not parent_raw.isdigit():
        return keep_parent_id, None, new_creds

    if parent_email:
        default_name = parent_name or f"Parent of {student_name}"
        user, plain, err = resolve_or_create_parent(
            school_id, name=default_name, email=parent_email
        )
        if err:
            return None, err, new_creds
        if user and plain:
            new_creds.append((user.name, user.email, plain))
        return user.id if user else None, None, new_creds

    if parent_raw.isdigit():
        parent_id = int(parent_raw)
        parent_user = User.query.get(parent_id)
        if (
            not parent_user
            or parent_user.role != "parent"
            or parent_user.school_id != school_id
        ):
            return None, "Selected parent must be a parent user at your school.", new_creds
        return parent_id, None, new_creds

    return None, None, new_creds


def link_parent_from_import_row(school_id, row: dict, cache: dict[str, int]):
    """Resolve parent for Excel row; cache maps email -> user id within one import."""
    email = (row.get("parent_email") or "").strip().lower()
    if not email:
        return None, None, None

    if email in cache:
        return cache[email], None, None

    name = (row.get("parent_name") or "").strip() or f"Parent of {row.get('name', 'learner')}"
    user, plain, err = resolve_or_create_parent(school_id, name=name, email=email)
    if err:
        return None, err, None
    if user:
        cache[email] = user.id
    cred = (user.name, user.email, plain) if user and plain else None
    return user.id if user else None, None, cred


# =======================
# STUDENTS (ADMIN, UG CLASS)
# =======================
def _student_status_map(student_ids):
    """Return {student_id: 'absent'|'present'} based on open Movement records."""
    if not student_ids:
        return {}
    open_movements = Movement.query.filter(
        Movement.student_id.in_(student_ids),
        Movement.time_in.is_(None),
        Movement.time_out.isnot(None),
    ).all()
    absent_ids = {m.student_id for m in open_movements}
    return {sid: ("absent" if sid in absent_ids else "present") for sid in student_ids}


@app.route("/students", methods=["GET"])
@login_required
def students():
    if current_user.role not in ("admin", "teacher") or not current_user.school_id:
        flash("Only school administrators and teachers can view learners.", "error")
        return redirect(url_for("dashboard"))

    filter_class = (request.args.get("klass") or "").strip()
    is_teacher = current_user.role == "teacher"

    query = Student.query.filter_by(school_id=current_user.school_id)
    if is_teacher and current_user.class_level:
        query = query.filter_by(class_level=current_user.class_level)
        if current_user.stream:
            query = query.filter_by(stream=current_user.stream)
        filter_class = current_user.class_level
    elif filter_class in UG_CLASS_CODES:
        query = query.filter_by(class_level=filter_class)
    student_rows = query.order_by(Student.class_level, Student.stream, Student.name).all()

    status_map = _student_status_map([s.id for s in student_rows])

    parents = (
        User.query.filter_by(role="parent", school_id=current_user.school_id)
        .order_by(User.name)
        .all()
    )
    parent_by_id = {p.id: p for p in parents}

    return render_template(
        "students.html",
        students=student_rows,
        parent_by_id=parent_by_id,
        status_map=status_map,
        ug_primary_levels=UG_PRIMARY_LEVELS,
        ug_secondary_levels=UG_SECONDARY_LEVELS,
        ug_class_codes=sorted(UG_CLASS_CODES),
        filter_class=filter_class,
        is_teacher=is_teacher,
        teacher_stream=current_user.stream if is_teacher else None,
        label_for_code=label_for_code,
    )


@app.route("/students/register", methods=["GET", "POST"])
@login_required
def register_student():
    if current_user.role != "teacher" or not current_user.school_id:
        flash("Only class teachers can register learners. School administrators view the roster added by teachers.", "error")
        return redirect(url_for("dashboard"))

    is_teacher = True

    parents = (
        User.query.filter_by(role="parent", school_id=current_user.school_id)
        .order_by(User.name)
        .all()
    )

    if request.method == "POST":
        name = (request.form.get("name") or "").strip()
        class_level = (request.form.get("class_level") or "").strip()
        stream = (request.form.get("stream") or "").strip()
        rfid_raw = (request.form.get("rfid_tag") or "").strip()

        if is_teacher and current_user.class_level:
            class_level = current_user.class_level
            stream = current_user.stream or ""

        if not name or class_level not in UG_CLASS_CODES:
            flash("Learner name and a valid class (P1–P7, S1–S6) are required.", "error")
            return redirect(url_for("register_student"))

        if is_teacher and current_user.class_level and class_level != current_user.class_level:
            flash("You can only register learners in your assigned class.", "error")
            return redirect(url_for("register_student"))

        parent_id, parent_err, new_parent_creds = link_parent_for_learner(
            current_user.school_id, name
        )
        if parent_err:
            flash(parent_err, "error")
            return redirect(url_for("register_student"))

        if rfid_raw:
            existing_tag = Student.query.filter_by(rfid_tag=rfid_raw).first()
            if existing_tag:
                flash("That RFID tag is already assigned to another learner.", "error")
                return redirect(url_for("register_student"))

        learner = Student(
            name=name,
            class_level=class_level,
            stream=stream or None,
            rfid_tag=rfid_raw,
            school_id=current_user.school_id,
            parent_id=parent_id,
        )
        db.session.add(learner)
        db.session.flush()
        try:
            learner_code = assign_student_code(learner)
        except ValueError as exc:
            db.session.rollback()
            flash(str(exc), "error")
            return redirect(url_for("register_student"))

        photo_file = request.files.get("photo")
        if photo_file and photo_file.filename:
            photo_name, photo_err = save_student_photo_upload(photo_file, learner.id)
            if photo_err:
                db.session.rollback()
                flash(photo_err, "error")
                return redirect(url_for("register_student"))
            learner.photo_filename = photo_name

        db.session.commit()
        flash(
            f"{name} was added to {class_display(class_level, stream)} "
            f"with learner ID {learner_code}.",
            "success",
        )
        flash_new_parent_credentials(new_parent_creds)
        return redirect(url_for("students", klass=class_level))

    return render_template(
        "students_register.html",
        parents=parents,
        ug_primary_levels=UG_PRIMARY_LEVELS,
        ug_secondary_levels=UG_SECONDARY_LEVELS,
        is_teacher=is_teacher,
        teacher_class=current_user.class_level if is_teacher else None,
        teacher_stream=current_user.stream if is_teacher else None,
    )


@app.route("/students/import/template")
@login_required
def student_import_template():
    if current_user.role != "teacher" or not current_user.school_id:
        flash("Only class teachers can download learner import templates.", "error")
        return redirect(url_for("dashboard"))
    buf = build_learner_import_template()
    buf.seek(0)
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="campuspass_learners_import_template.xlsx",
    )


@app.route("/students/import", methods=["POST"])
@login_required
def student_import_upload():
    if current_user.role != "teacher" or not current_user.school_id:
        flash("Only class teachers can import learners.", "error")
        return redirect(url_for("dashboard"))

    is_teacher = True

    f = request.files.get("file")
    if not f or not f.filename:
        flash("Choose an Excel file (.xlsx) to upload.", "error")
        return redirect(url_for("register_student"))

    if not f.filename.lower().endswith(".xlsx"):
        flash("Please upload an Excel workbook in .xlsx format.", "error")
        return redirect(url_for("register_student"))

    raw = f.read()

    rows, parse_errors = parse_learner_rows(raw)
    rows, dup_errors = filter_duplicate_rfids_in_file(rows)
    all_notes = parse_errors + dup_errors
    parent_cache: dict[str, int] = {}
    import_parent_creds: list[tuple[str, str, str]] = []
    cred_emails_shown: set[str] = set()

    added = 0
    for row in rows:
        row_num = row.pop("_row", 0)

        if is_teacher and current_user.class_level:
            if row["class_level"] != current_user.class_level:
                all_notes.append(
                    f"Row {row_num} ({row['name']!r}): class {row['class_level']} "
                    f"does not match your assigned class {current_user.class_level}, skipped."
                )
                continue

        row_stream = (row.get("stream") or "").strip()
        if is_teacher and current_user.stream:
            row_stream = current_user.stream

        tag = (row.get("rfid_tag") or "").strip()
        if tag:
            if Student.query.filter_by(rfid_tag=tag).first():
                all_notes.append(f"Row {row_num}: RFID {tag!r} already in the system, skipped.")
                continue

        parent_id, p_err, p_cred = link_parent_from_import_row(
            current_user.school_id, row, parent_cache
        )
        if p_err:
            all_notes.append(f"Row {row_num} ({row['name']!r}): {p_err} — learner imported without parent.")
        elif p_cred and p_cred[1] not in cred_emails_shown:
            import_parent_creds.append(p_cred)
            cred_emails_shown.add(p_cred[1])

        learner = Student(
            name=row["name"],
            class_level=row["class_level"],
            stream=row_stream or None,
            rfid_tag=row.get("rfid_tag") or "",
            school_id=current_user.school_id,
            parent_id=parent_id,
        )
        db.session.add(learner)
        try:
            db.session.flush()
            assign_student_code(learner)
        except ValueError as exc:
            db.session.rollback()
            flash(str(exc), "error")
            return redirect(url_for("register_student"))
        added += 1

    if added:
        db.session.commit()
        flash(f"Imported {added} learner(s) from the spreadsheet.", "success")
        if import_parent_creds:
            flash_parent_logins_download(import_parent_creds)
        if all_notes:
            body = "<br/>".join(html.escape(n) for n in all_notes[:35])
            if len(all_notes) > 35:
                body += f"<br/><em>…and {len(all_notes) - 35} more note(s).</em>"
            flash(Markup(f"<strong>Import notes</strong><br/>{body}"), "warning")
    else:
        db.session.rollback()
        if not rows and not all_notes:
            flash("No data rows were found to import.", "error")
        elif all_notes:
            body = "<br/>".join(html.escape(n) for n in all_notes[:35])
            if len(all_notes) > 35:
                body += f"<br/><em>…and {len(all_notes) - 35} more note(s).</em>"
            flash(
                Markup(
                    "<strong>No learners were imported.</strong><br/><br/>"
                    "<strong>Details</strong><br/>" + body
                ),
                "error",
            )
        else:
            flash("No learners were imported.", "error")

    return redirect(url_for("students"))


@app.route("/students/import/parent-logins")
@login_required
def download_parent_import_logins():
    if current_user.role != "teacher":
        flash("Only class teachers can download parent login spreadsheets.", "error")
        return redirect(url_for("dashboard"))

    rows = get_import_parent_logins()
    if not rows:
        flash("No parent login spreadsheet is available. Import learners with parent emails first.", "error")
        return redirect(url_for("register_student"))

    credentials = [(r["name"], r["email"], r["password"]) for r in rows]
    buf = build_parent_logins_workbook(credentials)
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="campuspass_parent_logins.xlsx",
    )


@app.route("/student-photos/<path:filename>")
@login_required
def student_photo(filename):
    safe = os.path.basename(filename)
    if not safe or safe != filename:
        abort(404)
    learner = Student.query.filter_by(photo_filename=safe).first_or_404()
    if not can_view_student_photo(current_user, learner):
        abort(403)
    return send_from_directory(student_photo_dir(), safe)


@app.route("/students/<int:student_id>/edit", methods=["GET", "POST"])
@login_required
def edit_student(student_id):
    learner = Student.query.get_or_404(student_id)
    if not can_manage_student(current_user, learner):
        flash("You do not have permission to edit this learner.", "error")
        return redirect(url_for("students"))

    parents = (
        User.query.filter_by(role="parent", school_id=current_user.school_id)
        .order_by(User.name)
        .all()
    )
    is_teacher = current_user.role == "teacher"

    if request.method == "POST":
        name = (request.form.get("name") or "").strip()
        class_level = (request.form.get("class_level") or "").strip().upper()
        stream = (request.form.get("stream") or "").strip()
        rfid_raw = (request.form.get("rfid_tag") or "").strip()

        if is_teacher and current_user.class_level:
            class_level = current_user.class_level
            stream = current_user.stream or ""

        if not name or class_level not in UG_CLASS_CODES:
            flash("Learner name and a valid class are required.", "error")
            return redirect(url_for("edit_student", student_id=student_id))

        parent_id, parent_err, new_parent_creds = link_parent_for_learner(
            current_user.school_id, name, keep_parent_id=learner.parent_id
        )
        if parent_err:
            flash(parent_err, "error")
            return redirect(url_for("edit_student", student_id=student_id))

        if rfid_raw:
            clash = Student.query.filter(
                Student.rfid_tag == rfid_raw, Student.id != learner.id
            ).first()
            if clash:
                flash("That RFID tag is already assigned to another learner.", "error")
                return redirect(url_for("edit_student", student_id=student_id))

        learner.name = name
        learner.class_level = class_level
        learner.stream = stream or None
        learner.rfid_tag = rfid_raw
        learner.parent_id = parent_id

        if is_teacher:
            if request.form.get("remove_photo") == "1" and learner.photo_filename:
                delete_student_photo_file(learner.photo_filename)
                learner.photo_filename = None
            else:
                photo_file = request.files.get("photo")
                if photo_file and photo_file.filename:
                    old_photo = learner.photo_filename
                    photo_name, photo_err = save_student_photo_upload(
                        photo_file, learner.id
                    )
                    if photo_err:
                        flash(photo_err, "error")
                        return redirect(url_for("edit_student", student_id=student_id))
                    learner.photo_filename = photo_name
                    if old_photo and old_photo != photo_name:
                        delete_student_photo_file(old_photo)

        db.session.commit()
        flash(f"{name} was updated.", "success")
        flash_new_parent_credentials(new_parent_creds)
        return redirect(url_for("students", klass=class_level))

    linked_parent = User.query.get(learner.parent_id) if learner.parent_id else None

    return render_template(
        "students_edit.html",
        learner=learner,
        linked_parent=linked_parent,
        parents=parents,
        ug_primary_levels=UG_PRIMARY_LEVELS,
        ug_secondary_levels=UG_SECONDARY_LEVELS,
        is_teacher=is_teacher,
        teacher_class=current_user.class_level if is_teacher else None,
        teacher_stream=current_user.stream if is_teacher else None,
    )


@app.route("/students/<int:student_id>/delete", methods=["POST"])
@login_required
def delete_student(student_id):
    learner = Student.query.get_or_404(student_id)
    if not can_manage_student(current_user, learner):
        flash("You do not have permission to remove this learner.", "error")
        return redirect(url_for("students"))

    name = learner.name
    klass = learner.class_level
    if learner.photo_filename:
        delete_student_photo_file(learner.photo_filename)
    Movement.query.filter_by(student_id=learner.id).delete(synchronize_session=False)
    Request.query.filter_by(student_id=learner.id).delete(synchronize_session=False)
    db.session.delete(learner)
    db.session.commit()
    flash(f"{name} was removed from the roster.", "success")
    return redirect(url_for("students", klass=klass))


# =======================
# DASHBOARD
# =======================
@app.route("/dashboard")
@login_required
def dashboard():
    pending_schools = []
    approved_schools = []
    if current_user.role == "superadmin":
        pending_schools = School.query.filter_by(status="pending").order_by(School.id.desc()).all()
        approved_schools = School.query.filter_by(status="approved").order_by(School.name).all()
    requests_list = school_scoped_requests(current_user)
    return render_template(
        "dashboard.html",
        requests=requests_list,
        pending_schools=pending_schools,
        approved_schools=approved_schools,
    )


# =======================
# VIEW REQUESTS
# =======================
@app.route("/requests")
@login_required
def requests_page():
    all_requests = school_scoped_requests(current_user)
    sid_set = {r.student_id for r in all_requests}
    student_by_id = {}
    if sid_set:
        student_by_id = {
            s.id: s
            for s in Student.query.filter(Student.id.in_(sid_set)).all()
        }
    roster = roster_for_requester(current_user)
    roster_picker = (
        students_picker_payload(roster)
        if current_user.role == "nurse" and roster
        else None
    )
    return render_template(
        "requests.html",
        requests=all_requests,
        student_by_id=student_by_id,
        roster_students=roster,
        roster_picker=roster_picker,
        label_for_code=label_for_code,
    )


# =======================
# CREATE REQUEST
# =======================
@app.route('/create_request', methods=['POST'])
@login_required
def create_request():
    if current_user.role not in ("nurse", "teacher", "parent"):
        flash("Your role cannot create movement requests.", "error")
        return redirect(url_for("dashboard"))

    student_id_raw = request.form.get("student_id")
    reason = (request.form.get("reason") or "").strip()

    if not student_id_raw or not reason:
        flash("Student and reason are required.", "error")
        return redirect(url_for("requests_page"))

    try:
        student_pk = int(student_id_raw)
    except (TypeError, ValueError):
        flash("Invalid student selection.", "error")
        return redirect(url_for("requests_page"))

    student = Student.query.get(student_pk)
    if not student:
        flash("Learner not found.", "error")
        return redirect(url_for("requests_page"))

    if current_user.role == "parent":
        if student.parent_id != current_user.id:
            flash("You can only request movement for your own children.", "error")
            return redirect(url_for("requests_page"))
    elif not student.school_id or student.school_id != current_user.school_id:
        flash("You can only request movement for learners at your school.", "error")
        return redirect(url_for("requests_page"))
    elif current_user.role == "teacher":
        if student.class_level != current_user.class_level:
            flash("You can only request movement for learners in your class.", "error")
            return redirect(url_for("requests_page"))
        if current_user.stream and student.stream != current_user.stream:
            flash("You can only request movement for learners in your stream.", "error")
            return redirect(url_for("requests_page"))

    req = Request(
        student_id=student_pk,
        requested_by=current_user.id,
        reason=reason,
        status="pending",
    )

    db.session.add(req)
    db.session.flush()

    cls = class_display(student.class_level, student.stream)
    notify_school_admins(
        student.school_id,
        f"New exit request for {student.name} ({cls}): {reason}. Request #{req.id}.",
        email_subject="CampusPass: new exit request",
    )
    db.session.commit()

    flash("Request submitted for administrator review.", "success")
    return redirect(url_for("requests_page"))


# =======================
# APPROVE / REJECT REQUEST
# =======================
@app.route('/handle_request/<int:id>', methods=['POST'])
@login_required
def handle_request(id):

    if current_user.role != "admin":
        flash("Only school administrators can approve requests.", "error")
        return redirect(url_for("requests_page"))

    decision = request.form.get("decision")

    if decision not in ["approved", "rejected"]:
        flash("Invalid approval action.", "error")
        return redirect(url_for("requests_page"))

    req = Request.query.get(id)

    if not req:
        flash("That request no longer exists.", "error")
        return redirect(url_for("requests_page"))

    req.status = decision
    req.approved_by = current_user.id

    student = Student.query.get(req.student_id)
    if student:
        cls = class_display(student.class_level, student.stream)
        msg = f"Exit request for {student.name} ({cls}) was {decision}. Reason: {req.reason}"
        if decision == "approved":
            notify_class_teacher(
                student,
                f"Exit request approved for {student.name} ({cls}). Reason: {req.reason}",
                kind=KIND_REQUEST_DECISION,
            )
        parent = User.query.get(student.parent_id) if student.parent_id else None
        if parent:
            alert_user(
                parent,
                msg,
                email_subject=f"CampusPass: request {decision} for {student.name}",
                kind=KIND_REQUEST_DECISION,
            )
        requester = User.query.get(req.requested_by) if req.requested_by else None
        if requester and requester.id != current_user.id:
            alert_user(
                requester,
                msg,
                email_subject=f"CampusPass: request {decision} for {student.name}",
                kind=KIND_REQUEST_DECISION,
            )

    db.session.commit()

    label = "approved" if decision == "approved" else "rejected"
    flash(f"Request #{req.id} was {label}.", "success")
    return redirect(url_for("requests_page"))


# =======================
# DOWNLOAD REQUESTS (EXCEL)
# =======================
@app.route("/requests/download")
@login_required
def download_requests():
    all_requests = school_scoped_requests(current_user)
    sid_set = {r.student_id for r in all_requests}
    student_by_id = {}
    if sid_set:
        student_by_id = {s.id: s for s in Student.query.filter(Student.id.in_(sid_set)).all()}
    user_by_id = {}
    uid_set = {r.requested_by for r in all_requests if r.requested_by}
    uid_set |= {r.approved_by for r in all_requests if r.approved_by}
    if uid_set:
        user_by_id = {u.id: u for u in User.query.filter(User.id.in_(uid_set)).all()}

    wb = Workbook()
    ws = wb.active
    ws.title = "Requests"
    headers = ["ID", "Student", "Class", "Stream", "Reason", "Status", "Requested by", "Approved by", "Date"]
    bold = Font(bold=True)
    for i, h in enumerate(headers, 1):
        ws.cell(row=1, column=i, value=h).font = bold
    for idx, r in enumerate(all_requests, 2):
        st = student_by_id.get(r.student_id)
        ws.cell(row=idx, column=1, value=r.id)
        ws.cell(row=idx, column=2, value=st.name if st else f"#{r.student_id}")
        ws.cell(row=idx, column=3, value=st.class_level if st else "")
        ws.cell(row=idx, column=4, value=st.stream if st and st.stream else "")
        ws.cell(row=idx, column=5, value=r.reason)
        ws.cell(row=idx, column=6, value=r.status)
        ws.cell(row=idx, column=7, value=user_by_id[r.requested_by].name if r.requested_by and r.requested_by in user_by_id else "")
        ws.cell(row=idx, column=8, value=user_by_id[r.approved_by].name if r.approved_by and r.approved_by in user_by_id else "")
        ws.cell(row=idx, column=9, value=r.created_at.strftime("%Y-%m-%d %H:%M") if r.created_at else "")
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name="campuspass_requests.xlsx")


# =======================
# CLEAR RESOLVED REQUESTS
# =======================
@app.route("/requests/clear", methods=["POST"])
@login_required
def clear_requests():
    if current_user.role != "admin":
        flash("Only administrators can clear requests.", "error")
        return redirect(url_for("requests_page"))
    cleared = Request.query.filter(
        Request.status.in_(["approved", "rejected", "completed"]),
        Request.student_id.in_(
            db.session.query(Student.id).filter_by(school_id=current_user.school_id)
        ),
    ).delete(synchronize_session=False)
    db.session.commit()
    flash(f"Cleared {cleared} resolved request(s).", "success")
    return redirect(url_for("requests_page"))


# =======================
# DOWNLOAD ROSTER (EXCEL)
# =======================
@app.route("/students/download")
@login_required
def download_roster():
    if current_user.role not in ("admin", "teacher") or not current_user.school_id:
        flash("You do not have permission to download the roster.", "error")
        return redirect(url_for("dashboard"))

    query = Student.query.filter_by(school_id=current_user.school_id)
    if current_user.role == "teacher" and current_user.class_level:
        query = query.filter_by(class_level=current_user.class_level)
        if current_user.stream:
            query = query.filter_by(stream=current_user.stream)
    filter_class = (request.args.get("klass") or "").strip()
    if filter_class in UG_CLASS_CODES and current_user.role == "admin":
        query = query.filter_by(class_level=filter_class)
    student_rows = query.order_by(Student.class_level, Student.stream, Student.name).all()
    status_map = _student_status_map([s.id for s in student_rows])

    parents = {p.id: p for p in User.query.filter_by(role="parent", school_id=current_user.school_id).all()}

    wb = Workbook()
    ws = wb.active
    ws.title = "Roster"
    headers = ["Name", "Class", "Stream", "Status", "RFID", "Parent", "Learner ID"]
    bold = Font(bold=True)
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    for i, h in enumerate(headers, 1):
        ws.cell(row=1, column=i, value=h).font = bold
    for idx, s in enumerate(student_rows, 2):
        ws.cell(row=idx, column=1, value=s.name)
        ws.cell(row=idx, column=2, value=s.class_level)
        ws.cell(row=idx, column=3, value=s.stream or "")
        status = status_map.get(s.id, "present")
        cell = ws.cell(row=idx, column=4, value=status.upper())
        cell.fill = green_fill if status == "present" else red_fill
        ws.cell(row=idx, column=5, value=s.rfid_tag or "")
        p = parents.get(s.parent_id)
        ws.cell(row=idx, column=6, value=p.name if p else "")
        ws.cell(row=idx, column=7, value=display_student_id(s))
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    suffix = f"_{filter_class}" if filter_class else ""
    return send_file(buf, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name=f"campuspass_roster{suffix}.xlsx")


# =======================
# NOTIFICATIONS
# =======================
@app.route("/notifications")
@login_required
def notifications():
    notes = Notification.query.filter_by(user_id=current_user.id).order_by(Notification.id.desc()).limit(100).all()
    return render_template("notifications.html", notifications=notes)


@app.route("/notifications/read", methods=["POST"])
@login_required
def mark_notifications_read():
    Notification.query.filter_by(user_id=current_user.id, is_read=False).update({"is_read": True})
    db.session.commit()
    flash("All notifications marked as read.", "success")
    return redirect(url_for("notifications"))


def _popup_kinds_for_role(role: str) -> list[str] | None:
    if role == "superadmin":
        return [KIND_SCHOOL_PENDING]
    if role == "admin":
        return [KIND_EXIT_REQUEST]
    if role == "teacher":
        return [KIND_DEPARTURE, KIND_RETURN]
    if role == "parent":
        return [KIND_DEPARTURE, KIND_RETURN, KIND_REQUEST_DECISION]
    return None


@app.route("/api/notifications/poll")
@login_required
def poll_notifications():
    """Return new notifications for live toast popups."""
    kinds = _popup_kinds_for_role(current_user.role)
    if not kinds:
        return jsonify({"notifications": [], "unread_count": 0})

    since_id = request.args.get("since_id", 0, type=int)
    q = Notification.query.filter(
        Notification.user_id == current_user.id,
        Notification.id > since_id,
        Notification.kind.in_(kinds),
    ).order_by(Notification.id.asc())

    items = []
    for n in q.limit(15).all():
        items.append(
            {
                "id": n.id,
                "message": n.message,
                "kind": n.kind or "general",
                "created_at": n.created_at.strftime("%H:%M") if n.created_at else "",
            }
        )

    unread = Notification.query.filter_by(
        user_id=current_user.id, is_read=False
    ).count()
    return jsonify({"notifications": items, "unread_count": unread})


# =======================
# SUPERADMIN APPROVES SCHOOL
# =======================
@app.route('/approve_school/<int:id>', methods=['POST'])
@login_required
def approve_school(id):

    if current_user.role != "superadmin":
        flash("You do not have permission to approve schools.", "error")
        return redirect(url_for("dashboard"))

    school = School.query.get(id)

    if not school:
        flash("School not found.", "error")
        return redirect(url_for("dashboard"))

    school.status = "approved"
    db.session.commit()

    flash(f"{school.name} is now approved. Administrators at that campus can sign in.", "success")
    return redirect(url_for("dashboard"))


# =======================
# SUPERADMIN DELETES SCHOOL
# =======================
@app.route('/delete_school/<int:id>', methods=['POST'])
@login_required
def delete_school(id):
    if current_user.role != "superadmin":
        flash("You do not have permission to delete schools.", "error")
        return redirect(url_for("dashboard"))

    school = School.query.get(id)
    if not school:
        flash("School not found.", "error")
        return redirect(url_for("dashboard"))

    name = school.name
    student_ids = [s.id for s in Student.query.filter_by(school_id=id).all()]
    if student_ids:
        Movement.query.filter(Movement.student_id.in_(student_ids)).delete(synchronize_session=False)
        Request.query.filter(Request.student_id.in_(student_ids)).delete(synchronize_session=False)
        Student.query.filter_by(school_id=id).delete(synchronize_session=False)
    user_ids = [u.id for u in User.query.filter_by(school_id=id).all()]
    if user_ids:
        Notification.query.filter(Notification.user_id.in_(user_ids)).delete(synchronize_session=False)
    User.query.filter_by(school_id=id).delete(synchronize_session=False)
    db.session.delete(school)
    db.session.commit()

    flash(f"School \"{name}\" and all its associated data have been deleted.", "success")
    return redirect(url_for("dashboard"))


# =======================
# GUARD PANEL (MANUAL FALLBACK)
# =======================
@app.route("/guard", methods=["GET"])
@login_required
def guard_panel():
    if current_user.role != "guard":
        flash("Only guards can access the gate panel.", "error")
        return redirect(url_for("dashboard"))
    return render_template("guard.html", student=None, search_done=False)


@app.route("/guard/lookup", methods=["POST"])
@login_required
def guard_lookup():
    if current_user.role != "guard":
        flash("Only guards can access the gate panel.", "error")
        return redirect(url_for("dashboard"))

    query = (request.form.get("query") or "").strip()
    student = find_student_by_lookup(
        query, school_id=current_user.school_id
    )

    if not student:
        flash(
            f"No learner found for \"{query}\". Try their learner ID (e.g. SKT-001), "
            "RFID tag, or name.",
            "error",
        )
        return render_template("guard.html", student=None, search_done=True)

    approved_req = Request.query.filter_by(student_id=student.id, status="approved").first()
    open_movement = Movement.query.filter_by(student_id=student.id, time_in=None).filter(
        Movement.time_out.isnot(None)
    ).first()
    is_out = open_movement is not None
    school = School.query.get(student.school_id) if student.school_id else None

    return render_template(
        "guard.html",
        student=student,
        school=school,
        approved_req=approved_req,
        open_movement=open_movement,
        is_out=is_out,
        search_done=True,
    )


@app.route("/guard/manual_scan/<int:student_id>", methods=["POST"])
@login_required
def guard_manual_scan(student_id):
    if current_user.role != "guard":
        flash("Only guards can use the gate panel.", "error")
        return redirect(url_for("dashboard"))

    student = Student.query.get(student_id)
    if not student:
        flash("Learner not found.", "error")
        return redirect(url_for("guard_panel"))

    open_movement = Movement.query.filter_by(student_id=student.id, time_in=None).filter(
        Movement.time_out.isnot(None)
    ).first()

    if open_movement:
        open_movement.time_in = datetime.utcnow()
        _complete_approved_exit_requests(student.id)
        notify_gate_movement(
            student,
            signed_in=True,
            time_label=datetime.utcnow().strftime("%H:%M"),
            manual_note=" (manual check-in by guard)",
        )
        db.session.commit()
        flash(f"{student.name} signed IN successfully. The exit approval has been closed.", "success")
        return redirect(url_for("guard_panel"))

    approved_req = Request.query.filter_by(student_id=student.id, status="approved").first()
    if not approved_req:
        flash(f"No approved exit request for {student.name}. Cannot sign out.", "error")
        return redirect(url_for("guard_panel"))

    move = Movement(
        student_id=student.id,
        time_out=datetime.utcnow(),
        guard_id=current_user.id,
    )
    db.session.add(move)
    notify_gate_movement(
        student,
        signed_in=False,
        time_label=datetime.utcnow().strftime("%H:%M"),
        reason=approved_req.reason,
        manual_note=" (manual by guard)",
    )
    db.session.commit()
    flash(f"{student.name} signed OUT successfully.", "success")
    return redirect(url_for("guard_panel"))


def _complete_approved_exit_requests(student_id: int) -> None:
    """Mark approved exit requests as completed when the learner returns to campus."""
    Request.query.filter_by(student_id=student_id, status="approved").update(
        {"status": "completed"},
        synchronize_session=False,
    )


def _student_is_signed_out(student_id: int) -> bool:
    return (
        Movement.query.filter_by(student_id=student_id, time_in=None)
        .filter(Movement.time_out.isnot(None))
        .first()
        is not None
    )


def _gate_scan_student_payload(
    student: Student,
    *,
    scan_action: str | None = None,
    exit_reason: str | None = None,
) -> dict:
    """Learner profile JSON for gate RFID display."""
    school = School.query.get(student.school_id) if student.school_id else None
    signed_out = _student_is_signed_out(student.id)
    return {
        "id": student.id,
        "name": student.name or "",
        "student_code": display_student_id(student),
        "class_level": student.class_level or "",
        "stream": student.stream or "",
        "class_label": class_display(student.class_level, student.stream),
        "rfid_tag": student.rfid_tag or "",
        "photo_url": student_photo_url(student) if student.photo_filename else None,
        "school_name": school.name if school else "",
        "on_campus": not signed_out,
        "scan_action": scan_action,
        "exit_reason": exit_reason or "",
    }


def _rfid_api_authorized():
    key = current_app.config.get("RFID_API_KEY") or ""
    if not key:
        return True
    supplied = request.headers.get("X-RFID-API-Key") or request.args.get("api_key")
    if request.is_json and request.get_json(silent=True):
        supplied = supplied or (request.get_json(silent=True) or {}).get("api_key")
    return supplied == key


def process_rfid_scan(tag: str, guard_id=None):
    """Shared RFID in/out logic for API and guard web scanner."""
    tag = (tag or "").strip()
    if not tag:
        return {"error": "RFID required"}, 400

    student = Student.query.filter_by(rfid_tag=tag).first()
    if not student:
        return {"error": "Invalid RFID — no learner is registered with this tag."}, 404

    if _student_is_signed_out(student.id):
        existing = Movement.query.filter_by(student_id=student.id, time_in=None).filter(
            Movement.time_out.isnot(None)
        ).first()
        existing.time_in = datetime.utcnow()
        _complete_approved_exit_requests(student.id)
        notify_gate_movement(
            student,
            signed_in=True,
            time_label=datetime.utcnow().strftime("%H:%M"),
        )
        db.session.commit()
        return {
            "message": "Student signed IN",
            "ok": True,
            "student": _gate_scan_student_payload(student, scan_action="in"),
        }, 200

    req = Request.query.filter_by(student_id=student.id, status="approved").first()
    if not req:
        return {
            "error": "No approved exit request for this learner.",
            "ok": False,
            "student": _gate_scan_student_payload(student, scan_action="denied"),
        }, 403

    move = Movement(
        student_id=student.id,
        time_out=datetime.utcnow(),
        guard_id=guard_id,
    )
    db.session.add(move)
    notify_gate_movement(
        student,
        signed_in=False,
        time_label=datetime.utcnow().strftime("%H:%M"),
        reason=req.reason,
    )
    db.session.commit()
    return {
        "message": "Student signed OUT",
        "ok": True,
        "student": _gate_scan_student_payload(
            student, scan_action="out", exit_reason=req.reason
        ),
    }, 200


# =======================
# RFID SCAN (IN + OUT) — API
# =======================
@app.route("/rfid_scan", methods=["POST"])
def rfid_scan():
    if not _rfid_api_authorized():
        return jsonify({"error": "Unauthorized"}), 401
    try:
        data = request.get_json(silent=True) or {}
        tag = data.get("rfid")
        body, code = process_rfid_scan(tag, guard_id=data.get("guard_id"))
        return jsonify(body), code
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/guard/scanner", methods=["GET"])
@login_required
def guard_scanner():
    if current_user.role != "guard":
        flash("Only guards can use the RFID scanner.", "error")
        return redirect(url_for("dashboard"))
    return render_template("guard_scanner.html")


@app.route("/guard/scanner/scan", methods=["POST"])
@login_required
def guard_scanner_scan():
    if current_user.role != "guard":
        return jsonify({"error": "Forbidden"}), 403
    data = request.get_json(silent=True) or {}
    tag = data.get("rfid") or request.form.get("rfid")
    body, code = process_rfid_scan(tag, guard_id=current_user.id)
    return jsonify(body), code


# =======================
# CHANGE PASSWORD (ALL USERS)
# =======================
@app.route("/account/password", methods=["GET", "POST"])
@login_required
def change_password():
    if request.method == "POST":
        current_pw = request.form.get("current_password") or ""
        new_pw = request.form.get("new_password") or ""
        confirm_pw = request.form.get("confirm_password") or ""

        if not password_matches(current_user.password, current_pw):
            flash("Your current password is incorrect.", "error")
            return redirect(url_for("change_password"))

        if len(new_pw) < MIN_PASSWORD_LENGTH:
            flash(
                f"New password must be at least {MIN_PASSWORD_LENGTH} characters.",
                "error",
            )
            return redirect(url_for("change_password"))

        if new_pw != confirm_pw:
            flash("New password and confirmation do not match.", "error")
            return redirect(url_for("change_password"))

        if password_matches(current_user.password, new_pw):
            flash("Choose a password that is different from your current one.", "error")
            return redirect(url_for("change_password"))

        current_user.password = hash_password(new_pw)
        db.session.commit()
        flash("Your password was updated successfully.", "success")
        return redirect(url_for("dashboard"))

    return render_template("change_password.html")


# =======================
# LOGOUT
# =======================
@app.route('/logout')
@login_required
def logout():
    logout_user()
    return redirect(url_for("login"))


# =======================
# RUN APP + CREATE SUPERADMIN
# =======================
if __name__ == "__main__":
    with app.app_context():
        db.create_all()
        run_migrations()
        ensure_superadmin_exists()

    app.run(debug=True)