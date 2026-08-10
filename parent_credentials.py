"""Build and session-store parent login spreadsheets after bulk import."""

from __future__ import annotations

from io import BytesIO

from flask import flash, session, url_for
from markupsafe import Markup
from openpyxl import Workbook
from openpyxl.styles import Font

SESSION_KEY = "parent_import_logins"


def store_import_parent_logins(credentials: list[tuple[str, str, str]]) -> None:
    """Each item is (parent_name, email, plain_password)."""
    session[SESSION_KEY] = [
        {"name": name, "email": email, "password": password}
        for name, email, password in credentials
    ]


def clear_import_parent_logins() -> None:
    session.pop(SESSION_KEY, None)


def get_import_parent_logins() -> list[dict[str, str]]:
    rows = session.get(SESSION_KEY) or []
    if not isinstance(rows, list):
        return []
    return [r for r in rows if isinstance(r, dict) and r.get("email")]


def build_parent_logins_workbook(credentials: list[tuple[str, str, str]]) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Parent logins"
    bold = Font(bold=True)
    headers = ["Parent name", "Email (login)", "Temporary password"]
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = bold
    for row_idx, (name, email, password) in enumerate(credentials, start=2):
        ws.cell(row=row_idx, column=1, value=name)
        ws.cell(row=row_idx, column=2, value=email)
        ws.cell(row=row_idx, column=3, value=password)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 22
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def flash_parent_logins_download(credentials: list[tuple[str, str, str]]) -> None:
    if not credentials:
        return
    store_import_parent_logins(credentials)
    count = len(credentials)
    label = "login" if count == 1 else "logins"
    flash(
        Markup(
            f"<strong>{count} parent {label} created.</strong> "
            f'<a href="{url_for("download_parent_import_logins")}" class="alert-link fw-semibold">'
            "Download credentials spreadsheet (.xlsx)</a> "
            "and share each row securely with the family."
        ),
        "success",
    )
